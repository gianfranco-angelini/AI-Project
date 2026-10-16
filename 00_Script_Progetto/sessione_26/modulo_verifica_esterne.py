# Sessione 26 (16/10/2026) - SOLA LETTURA del database
# Crea il modulo Excel per gli Acquisti con le fasi esterne PRONTE della T09-0100:
# per ogni fase dire se il materiale e' gia' presso il fornitore, con data di uscita, n. ordine SAP e data promessa.
# Legge il CSV di prossime_fasi.py (rilanciarlo prima se il CSV non e' di oggi) e completa i dati dalle Job Card.
# Uso in bench --site site1.local console:
#   exec(open("/home/frappe-user/script_progetto/modulo_verifica_esterne.py").read(), {"frappe": frappe})
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

CSV_IN = "/home/frappe-user/script_progetto/prossime_fasi_t09.csv"
OUT = "/home/frappe-user/script_progetto/Verifica_Esterne_Avvio_T09.xlsx"

righe = []
for r in list(csv.reader(open(CSV_IN, encoding="utf-8-sig"), delimiter=";"))[1:]:
    if r[0] == "PRONTA" and r[6] == "Lavorazione Esterna":
        righe.append(r)
righe.sort(key=lambda x: (x[8], x[1], x[4]))
print("Fasi esterne pronte:", len(righe))

wb = Workbook()
ist = wb.active
ist.title = "Istruzioni"
testo = [
    "MODULO DI VERIFICA - LAVORAZIONI ESTERNE DA AVVIARE - T09-0100 (Job 26-19008)",
    "Estratto dal sistema M.E.S. Remazel il " + frappe.utils.formatdate(frappe.utils.today(), "dd/MM/yyyy") + ".",
    "",
    "Una riga per ogni fase esterna pronta da avviare (prime lavorazioni presso i fornitori, tutti i set).",
    "Le colonne grigie vengono dal sistema: non modificarle.",
    "Compilare solo le colonne gialle:",
    "  - Gia' presso il fornitore: Si = materiale gia' spedito; No = ancora da spedire",
    "  - Data uscita: data reale di spedizione (solo se Si)",
    "  - N. ordine SAP: numero dell'ordine di conto lavoro, se gia' emesso",
    "  - Data promessa: data di rientro confermata dal fornitore, se nota",
    "  - Note: qualsiasi informazione utile (rientro parziale, ordine unico per piu' set, ecc.)",
    "",
    "Le righe con Si verranno registrate come 'Presso fornitore' con la data di uscita indicata;",
    "le righe con No restano 'Da inviare' e si registrano alla spedizione (Guida operativa, capitolo 6.3).",
    "Il file compilato va restituito all'amministratore del sistema, che lo carica in un unico passaggio.",
]
for t in testo:
    ist.append([t])
ist["A1"].font = Font(bold=True, size=12)
ist.column_dimensions["A"].width = 110

ws = wb.create_sheet("Esterne")
grigie = ["Fornitore", "Consegna set", "Production Plan", "Work Order", "Articolo", "Job Card",
          "Descrizione fase", "Quantita'", "Inizio pianificato", "N. ordine SAP a sistema"]
gialle = ["Gia' presso il fornitore (Si/No)", "Data uscita", "N. ordine SAP", "Data promessa", "Note"]
ws.append(grigie + gialle)
grigio = PatternFill("solid", fgColor="D9D9D9")
giallo = PatternFill("solid", fgColor="FFF2CC")
testa = PatternFill("solid", fgColor="757F9B")
bordo = Border(*(Side(style="thin", color="BBBBBB"),) * 4)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = testa
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = bordo

for r in righe:
    jc = frappe.db.get_value("Job Card", r[5], ["for_quantity", "custom_ordine_sap"], as_dict=True) or {}
    ws.append([r[8], r[1], r[2], r[3], r[4], r[5], r[7], jc.get("for_quantity"), r[10][:10],
               jc.get("custom_ordine_sap") or "", "", "", "", "", ""])

n = ws.max_row
for riga in ws.iter_rows(min_row=2, max_row=n):
    for i, c in enumerate(riga):
        c.fill = grigio if i < len(grigie) else giallo
        c.border = bordo
larghezze = [24, 13, 19, 20, 17, 14, 60, 10, 15, 16, 16, 13, 16, 14, 40]
lettere = "ABCDEFGHIJKLMNO"
for i, w in enumerate(larghezze):
    ws.column_dimensions[lettere[i]].width = w
ws.row_dimensions[1].height = 32
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:O%d" % n
dv = DataValidation(type="list", formula1='"Si,No"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("K2:K%d" % n)

wb.save(OUT)
print("Creato:", OUT, "| righe:", n - 1)
