# Sessione 23 - SOLA LETTURA sul database. Crea il modulo Excel di rilevazione dello stato reale dei set 1 e 2 T09.
# Uso (console): exec(open("/home/frappe-user/script_progetto/modulo_rilevazione_set12.py").read(), {"frappe": frappe})
# Output: /home/frappe-user/script_progetto/Rilevazione_Set_1_2_T09.xlsx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT = "PROJ-0002"
OUT = "/home/frappe-user/script_progetto/Rilevazione_Set_1_2_T09.xlsx"

consegne = [r[0] for r in frappe.db.sql("""
    SELECT DISTINCT jc.custom_consegna_set FROM `tabJob Card` jc
    JOIN `tabWork Order` wo ON wo.name = jc.work_order
    WHERE wo.project = %s AND wo.docstatus = 1 AND jc.docstatus < 2 AND jc.custom_consegna_set IS NOT NULL
    ORDER BY jc.custom_consegna_set""", PROJECT)][:2]
print("Consegne set 1 e 2:", consegne)

righe = frappe.db.sql("""
    SELECT jc.custom_consegna_set AS consegna, jc.custom_production_plan AS piano, jc.work_order, jc.production_item,
           jc.sequence_id, jc.name AS job_card, jc.operation, jc.custom_descrizione_fase AS descrizione,
           IFNULL(jc.workstation, jc.workstation_type) AS reparto, jc.custom_fornitore AS fornitore,
           jc.status, jc.expected_start_date, jc.expected_end_date
    FROM `tabJob Card` jc JOIN `tabWork Order` wo ON wo.name = jc.work_order
    WHERE wo.project = %s AND wo.docstatus = 1 AND jc.docstatus < 2 AND jc.custom_consegna_set IN %s
    ORDER BY jc.custom_consegna_set, jc.work_order, jc.sequence_id, jc.name""", (PROJECT, tuple(consegne)), as_dict=True)
print("Job Card estratte:", len(righe))

wb = Workbook()
ist = wb.active
ist.title = "Istruzioni"
for i, t in enumerate([
    "MODULO DI RILEVAZIONE - SET 1 E 2 T09-0100",
    "",
    "Una riga per ogni fase (Job Card). Le colonne grigie vengono dal sistema: non modificarle.",
    "Compilare solo le colonne gialle:",
    "  Eseguita: Si = fase terminata / In corso = iniziata non finita / No = non iniziata",
    "  Data fine reale: per le fasi eseguite (gg/mm/aaaa); se non nota, data approssimativa",
    "  Quantita completata: se diversa da quella della fase",
    "  Per le fasi esterne In corso: Data uscita, N. ordine SAP, Data promessa (se nota)",
    "  Note: qualsiasi informazione utile",
    "Le righe lasciate vuote restano come pianificate dal sistema.",
], 1):
    ist.cell(row=i, column=1, value=t).font = Font(bold=(i == 1))
ist.column_dimensions["A"].width = 110

ws = wb.create_sheet("Rilevazione")
sistema = ["Consegna set", "Production Plan", "Work Order", "Articolo", "Seq.", "Job Card", "Operazione",
           "Descrizione fase", "Reparto", "Fornitore", "Stato a sistema", "Inizio pianificato", "Fine pianificata"]
utente = ["Eseguita", "Data fine reale", "Quantita completata", "Data uscita", "N. ordine SAP", "Data promessa", "Note"]
grigio = PatternFill("solid", fgColor="D9D9D9")
giallo = PatternFill("solid", fgColor="FFF2CC")
for c, t in enumerate(sistema + utente, 1):
    cell = ws.cell(row=1, column=c, value=t)
    cell.font = Font(bold=True)
    cell.fill = grigio if c <= len(sistema) else giallo
    cell.alignment = Alignment(wrap_text=True, vertical="top")
for r, x in enumerate(righe, 2):
    vals = [x.consegna, x.piano, x.work_order, x.production_item, x.sequence_id, x.job_card, x.operation,
            x.descrizione, x.reparto, x.fornitore, x.status,
            str(x.expected_start_date or "")[:16], str(x.expected_end_date or "")[:16]]
    for c, v in enumerate(vals, 1):
        ws.cell(row=r, column=c, value=v)
dv = DataValidation(type="list", formula1='"Si,In corso,No"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("N2:N%d" % (len(righe) + 1))
larg = [12, 18, 20, 22, 6, 16, 20, 50, 22, 28, 14, 17, 17, 10, 14, 12, 12, 16, 14, 30]
for i, w in enumerate(larg, 1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:T%d" % (len(righe) + 1)
wb.save(OUT)
print("Scritto", OUT)
